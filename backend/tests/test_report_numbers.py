import csv
import importlib.util
import sys
from pathlib import Path

import pytest

REPO = Path(__file__).resolve().parents[2]
spec = importlib.util.spec_from_file_location("report_numbers", REPO / "scripts" / "report_numbers.py")
report_numbers = importlib.util.module_from_spec(spec)
sys.modules["report_numbers"] = report_numbers
spec.loader.exec_module(report_numbers)


def _e1_rows():
    with (REPO / "experiments" / "E1_algorithm_comparison" / "results.csv").open(newline="") as f:
        return {r["algorithm_key"]: r for r in csv.DictReader(f)}


def test_headline_numbers_are_derived_from_the_e1_csv():
    rows = _e1_rows()
    greedy_km = float(rows["greedy"]["total_distance"])
    qpso_km = float(rows["qpso_ls"]["total_distance"])
    expected = round((greedy_km - qpso_km) / greedy_km * 100.0, 1)

    numbers = report_numbers.build_numbers()

    assert numbers["e1_headline"]["km_fewer_vs_greedy_pct"] == expected
    assert numbers["e1"]["qpso_ls"]["total_cost"] == float(rows["qpso_ls"]["total_cost"])


def test_real_city_ranges_come_from_the_validation_file():
    rc = report_numbers.build_numbers()["real_cities"]
    assert len(rc["cities"]) >= 1
    low, high = rc["qpso_below_greedy_pct_range"]
    assert low <= high
    assert all(low <= c["qpso_below_greedy_pct"] <= high for c in rc["cities"])


def test_find_missing_flags_a_stale_figure_and_passes_a_current_one():
    numbers = report_numbers.build_numbers()
    expected = report_numbers.expected_deck_strings(numbers)
    current = "\n".join(s for _, s in expected)
    stale = current.replace(f"{numbers['e1']['qpso_ls']['total_cost']:.2f}", "191.40")

    assert report_numbers.find_missing(current, expected) == []
    missing_labels = [label for label, _ in report_numbers.find_missing(stale, expected)]
    assert "qpso_ls cost" in missing_labels


def test_deck_check_reads_text_and_tables_from_a_real_pptx(tmp_path):
    pptx = pytest.importorskip("pptx")
    from pptx.util import Inches

    numbers = report_numbers.build_numbers()
    q = numbers["e1"]["qpso_ls"]

    prs = pptx.Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    table = slide.shapes.add_table(1, 1, Inches(1), Inches(1), Inches(3), Inches(1)).table
    table.cell(0, 0).text = f"{q['total_cost']:.2f}"
    path = tmp_path / "deck.pptx"
    prs.save(str(path))

    text = report_numbers.extract_pptx_text(path)

    assert f"{q['total_cost']:.2f}" in text
